from pathlib import Path
import json
from docx import Document
from pptx import Presentation
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
MANIFEST_PATH = ROOT / "data" / "manifests" / "classified_documents.json"
OUT_ROOT = ROOT / "data" / "extracted"


def extract_docx(path: Path):
    doc = Document(str(path))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    return {"paragraphs": paragraphs}


def extract_pptx(path: Path):
    prs = Presentation(str(path))
    slides = []
    for idx, slide in enumerate(prs.slides, start=1):
        texts = [shape.text for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()]
        slides.append({"slide": idx, "texts": texts})
    return {"slides": slides}


def extract_xlsx(path: Path):
    wb = load_workbook(filename=str(path), data_only=True)
    sheets = {}
    for ws in wb.worksheets:
        sheets[ws.title] = list(ws.iter_rows(values_only=True))
    return {"sheets": sheets}


def main():
    if not MANIFEST_PATH.exists():
        raise FileNotFoundError(f"Missing manifest: {MANIFEST_PATH}")

    OUT_ROOT.mkdir(parents=True, exist_ok=True)
    manifest = json.loads(MANIFEST_PATH.read_text(encoding="utf-8"))

    for item in manifest:
        source_path = ROOT / item["path"]
        kind = item["classification"]

        if kind == "docx":
            data = extract_docx(source_path)
        elif kind == "pptx":
            data = extract_pptx(source_path)
        elif kind in {"xlsx", "xls"}:
            data = extract_xlsx(source_path)
        else:
            continue

        out_path = OUT_ROOT / f"{Path(item['path']).stem}_office.json"
        out_path.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")
        print(f"Office extraction complete: {source_path}")


if __name__ == "__main__":
    main()
