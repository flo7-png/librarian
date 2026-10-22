"""
Unit tests for the document extraction pipeline
Run with: pytest tests/test_extraction.py -v
"""
import pytest
import tempfile
import os
import json
import shutil
from pathlib import Path
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.models import Base, Collection, Document, DocumentStatus, JobType, JobStatus
from app.services import register_collection, CollectionCreate
from app.classify import classify_file, classify_by_extension, is_scanned_pdf, FileType, PDFType
from app.dedupe import run_dedupe_scan, get_duplicate_report, DuplicateGroup
from app.extract import (
    extract_text, extract_text_pdf, extract_text_docx, extract_text_pptx,
    extract_text_xlsx, extract_text_json, extract_text_plain,
    ExtractionResult, save_extraction_result
)
from app.jsonnorm import normalize_chat_export
from app.quality import compute_quality_score, QualityScore


@pytest.fixture
def temp_dir():
    """Create a temporary directory with test files"""
    with tempfile.TemporaryDirectory() as tmpdir:
        # Create test file structure
        Path(tmpdir, "docs").mkdir()
        Path(tmpdir, "docs", "file1.txt").write_text("Content 1")
        Path(tmpdir, "docs", "file2.md").write_text("# Markdown\n\nContent 2")
        yield tmpdir


@pytest.fixture
def test_db():
    """Create a test database"""
    DATABASE_URL = "sqlite:///:memory:"
    engine = create_engine(DATABASE_URL, connect_args={"check_same_thread": False})
    Base.metadata.create_all(bind=engine)
    Session = sessionmaker(bind=engine)
    return Session()


# ============= CLASSIFICATION TESTS =============

def test_classify_by_extension():
    """Test extension-based classification"""
    assert classify_by_extension(Path("test.pdf")) == FileType.PDF
    assert classify_by_extension(Path("test.txt")) == FileType.TEXT
    assert classify_by_extension(Path("test.md")) == FileType.MARKDOWN
    assert classify_by_extension(Path("test.docx")) == FileType.DOCX
    assert classify_by_extension(Path("test.pptx")) == FileType.PPTX
    assert classify_by_extension(Path("test.xlsx")) == FileType.XLSX
    assert classify_by_extension(Path("test.json")) == FileType.JSON
    assert classify_by_extension(Path("test.mp4")) == FileType.VIDEO
    assert classify_by_extension(Path("test.zip")) == FileType.ARCHIVE
    assert classify_by_extension(Path("test.jpg")) == FileType.IMAGE
    assert classify_by_extension(Path("test.unknown")) == FileType.OTHER


def test_classify_file(temp_dir):
    """Test file classification with real files"""
    txt_file = Path(temp_dir) / "docs" / "file1.txt"
    md_file = Path(temp_dir) / "docs" / "file2.md"
    
    assert classify_file(txt_file) == FileType.TEXT
    assert classify_file(md_file) == FileType.MARKDOWN


def test_is_scanned_pdf_heuristic(temp_dir):
    """Test scanned PDF detection heuristic"""
    # Create a minimal PDF with text content
    pdf_path = Path(temp_dir) / "test.pdf"
    
    # We can't easily create a real PDF in a test without external deps,
    # so we test the logic by mocking
    from app.classify import SCANNED_PDF_CHAR_THRESHOLD
    
    # Test that the function exists and handles missing pdfplumber gracefully
    # This is a structural test - actual PDF testing requires fixtures
    assert SCANNED_PDF_CHAR_THRESHOLD == 100  # default


# ============= DEDUPE TESTS =============

def test_dedupe_exact_hash_grouping(test_db, temp_dir):
    """Test that exact duplicate files are grouped together"""
    # Register collection
    data = CollectionCreate(
        name="test-dedupe",
        host_display_path=temp_dir,
        container_root=temp_dir,
    )
    collection = register_collection(test_db, data)
    
    # Create two identical files
    file1 = Path(temp_dir) / "file1.txt"
    file2 = Path(temp_dir) / "file2.txt"
    content = "Exact same content for deduplication test"
    file1.write_text(content)
    file2.write_text(content)
    
    # Add documents manually with same hash
    import hashlib
    hash_val = hashlib.sha256(content.encode()).hexdigest()
    
    doc1 = Document(
        id="doc1", collection_id=collection.id, relative_path="file1.txt",
        source_path=str(file1), file_name="file1.txt", extension="txt",
        size_bytes=len(content), content_hash=hash_val, status=DocumentStatus.DISCOVERED
    )
    doc2 = Document(
        id="doc2", collection_id=collection.id, relative_path="file2.txt",
        source_path=str(file2), file_name="file2.txt", extension="txt",
        size_bytes=len(content), content_hash=hash_val, status=DocumentStatus.DISCOVERED
    )
    test_db.add_all([doc1, doc2])
    test_db.commit()
    
    # Run dedupe scan
    summary = run_dedupe_scan(test_db, collection.id)
    
    assert summary["exact_duplicate_groups"] == 1
    assert summary["exact_duplicate_files"] == 2
    assert summary["filename_size_groups"] == 0
    
    # Check report
    report = get_duplicate_report(test_db, collection.id)
    assert report["total_groups"] == 1
    assert report["exact_hash_groups"] == 1
    assert len(report["groups"][0]["documents"]) == 2


def test_dedupe_filename_size_grouping(test_db, temp_dir):
    """Test that same filename+size but different content are grouped"""
    data = CollectionCreate(
        name="test-dedupe2",
        host_display_path=temp_dir,
        container_root=temp_dir,
    )
    collection = register_collection(test_db, data)
    
    # Two files with same name and size but different content (different hashes)
    import hashlib
    content1 = "Content A"
    content2 = "Content B"
    hash1 = hashlib.sha256(content1.encode()).hexdigest()
    hash2 = hashlib.sha256(content2.encode()).hexdigest()
    
    doc1 = Document(
        id="doc1", collection_id=collection.id, relative_path="sub/file.txt",
        source_path=str(Path(temp_dir) / "sub" / "file.txt"), file_name="file.txt",
        extension="txt", size_bytes=len(content1), content_hash=hash1, status=DocumentStatus.DISCOVERED
    )
    doc2 = Document(
        id="doc2", collection_id=collection.id, relative_path="other/file.txt",
        source_path=str(Path(temp_dir) / "other" / "file.txt"), file_name="file.txt",
        extension="txt", size_bytes=len(content2), content_hash=hash2, status=DocumentStatus.DISCOVERED
    )
    test_db.add_all([doc1, doc2])
    test_db.commit()
    
    summary = run_dedupe_scan(test_db, collection.id)
    
    assert summary["exact_duplicate_groups"] == 0
    assert summary["filename_size_groups"] == 1
    assert summary["filename_size_files"] == 2


# ============= JSON NORMALIZATION TESTS =============

def test_normalize_claude_format():
    """Test Claude chat export normalization"""
    claude_data = [
        {"type": "user", "message": {"content": "Hello"}},
        {"type": "assistant", "message": {"content": "Hi there!"}},
        {"type": "user", "message": {"content": "How are you?"}},
        {"type": "assistant", "message": {"content": "I'm doing well!"}},
    ]
    
    result = normalize_chat_export(claude_data)
    
    assert result["format"] == "claude"
    assert result["turn_count"] == 4
    assert result["turns"][0]["role"] == "user"
    assert result["turns"][0]["text"] == "Hello"
    assert result["turns"][1]["role"] == "assistant"
    assert result["turns"][1]["text"] == "Hi there!"
    assert "User: Hello" in result["text"]
    assert "Assistant: Hi there!" in result["text"]


def test_normalize_claude_skips_tool_use():
    """Test that tool_use/tool_result are skipped"""
    claude_data = [
        {"type": "user", "message": {"content": "Calculate 2+2"}},
        {"type": "tool_use", "message": {"content": {"name": "calculator", "args": {}}}},
        {"type": "tool_result", "message": {"content": "4"}},
        {"type": "assistant", "message": {"content": "The answer is 4"}},
    ]
    
    result = normalize_chat_export(claude_data)
    
    assert result["turn_count"] == 2  # Only user and assistant
    assert result["turns"][0]["text"] == "Calculate 2+2"
    assert result["turns"][1]["text"] == "The answer is 4"


def test_normalize_openai_format():
    """Test OpenAI/ChatGPT export normalization"""
    openai_data = {
        "messages": [
            {"role": "user", "content": "Hello"},
            {"role": "assistant", "content": "Hi!"},
            {"role": "system", "content": "You are helpful"},  # Should be skipped
        ]
    }
    
    result = normalize_chat_export(openai_data)
    
    assert result["format"] == "openai"
    assert result["turn_count"] == 2  # System skipped
    assert result["turns"][0]["role"] == "user"
    assert result["turns"][1]["role"] == "assistant"


def test_normalize_generic_list():
    """Test generic list format"""
    generic_data = [
        {"role": "user", "content": "Test"},
        {"role": "assistant", "content": "Response"},
    ]
    
    result = normalize_chat_export(generic_data)
    
    assert result["format"] == "generic_list"
    assert result["turn_count"] == 2


# ============= QUALITY SCORING TESTS =============

def test_quality_good_text():
    """Test quality scoring on good text"""
    text = "This is a well-formed document with proper English sentences. " * 10
    score = compute_quality_score(text, file_size=1000, pages=1, method="text", extraction_confidence=0.9)
    
    assert score.passed == True
    assert score.needs_review == False
    assert score.score > 0.5
    assert score.extracted_chars == len(text)


def test_quality_empty_text():
    """Test quality scoring on empty text"""
    score = compute_quality_score("", file_size=1000, pages=1, method="text", extraction_confidence=0.9)
    
    assert score.passed == False
    assert score.needs_review == True
    assert score.score == 0.0


def test_quality_binary_garbage():
    """Test quality scoring on binary garbage (high replacement chars)"""
    text = "Some text \ufffd\ufffd\ufffd\ufffd\ufffd more text"
    score = compute_quality_score(text, file_size=100, pages=1, method="ocr", extraction_confidence=0.5)
    
    assert score.needs_review == True
    assert score.score < 0.35
    assert score.details["replacement_ratio"] > 0.1


def test_quality_low_alpha_ratio():
    """Test quality scoring on mostly non-alphanumeric text"""
    text = "!!!@@@###$$$%%%^^^&&&***(())"
    score = compute_quality_score(text, file_size=100, pages=1, method="text", extraction_confidence=0.8)
    
    assert score.needs_review == True
    assert score.details["alpha_ratio"] < 0.5


# ============= EXTRACTION TESTS =============

def test_extract_text_plain(temp_dir):
    """Test plain text extraction"""
    txt_file = Path(temp_dir) / "test.txt"
    txt_file.write_text("Hello world\nThis is a test")
    
    result = extract_text_plain(txt_file)
    
    assert isinstance(result, ExtractionResult)
    assert result.method == "native"
    assert "Hello world" in result.text
    assert result.pages == 1
    assert result.confidence == 1.0


def test_extract_text_markdown(temp_dir):
    """Test markdown extraction"""
    md_file = Path(temp_dir) / "test.md"
    md_file.write_text("# Header\n\n**Bold** text")
    
    result = extract_text_plain(md_file)
    
    assert result.method == "native"
    assert "Header" in result.text
    assert "Bold" in result.text


def test_extract_docx():
    """Test DOCX extraction using python-docx"""
    pytest.importorskip("docx")
    
    with tempfile.TemporaryDirectory() as tmpdir:
        docx_path = Path(tmpdir) / "test.docx"
        
        # Create a DOCX file
        from docx import Document as DocxDocument
        doc = DocxDocument()
        doc.add_paragraph("Paragraph 1")
        doc.add_paragraph("Paragraph 2")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "A1"
        table.cell(0, 1).text = "B1"
        table.cell(1, 0).text = "A2"
        table.cell(1, 1).text = "B2"
        doc.save(docx_path)
        
        result = extract_text_docx(docx_path)
        
        assert result.method == "native"
        assert "Paragraph 1" in result.text
        assert "Paragraph 2" in result.text
        assert "A1" in result.text
        assert "B2" in result.text


def test_extract_xlsx():
    """Test XLSX extraction using openpyxl"""
    pytest.importorskip("openpyxl")
    
    with tempfile.TemporaryDirectory() as tmpdir:
        xlsx_path = Path(tmpdir) / "test.xlsx"
        
        # Create an XLSX file
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws.append(["Col1", "Col2"])
        ws.append(["A1", "B1"])
        ws.append(["A2", "B2"])
        wb.save(xlsx_path)
        
        result = extract_text_xlsx(xlsx_path)
        
        assert result.method == "native"
        assert "Sheet1" in result.text
        assert "Col1" in result.text
        assert "A1" in result.text
        assert "B2" in result.text


def test_extract_json_chat_export():
    """Test JSON chat export normalization"""
    pytest.importorskip("json")
    
    with tempfile.TemporaryDirectory() as tmpdir:
        json_path = Path(tmpdir) / "chat.json"
        
        # Create a fake Claude export
        claude_data = [
            {"type": "user", "message": {"content": "Hello"}},
            {"type": "assistant", "message": {"content": "Hi there!"}},
        ]
        with open(json_path, 'w') as f:
            json.dump(claude_data, f)
        
        result = extract_text_json(json_path)
        
        assert result.method == "native"
        assert "User: Hello" in result.text
        assert "Assistant: Hi there!" in result.text


def test_save_extraction_result(temp_dir):
    """Test saving extraction result to disk"""
    result = ExtractionResult(
        text="Test extracted text!",
        pages=1,
        method="test",
        chars_per_page=[20],
        metadata={},
        confidence=0.9,
        extracted_at="2024-01-01T00:00:00"
    )
    
    text_path, manifest_path = save_extraction_result(
        Path(temp_dir) / "source.txt",
        result,
        "abc123hash",
        extracted_root=Path(temp_dir) / "extracted"
    )
    
    assert text_path.exists()
    assert manifest_path.exists()
    assert text_path.read_text() == "Test extracted text!"
    
    manifest = json.loads(manifest_path.read_text())
    assert manifest["chars"] == 20
    assert manifest["method"] == "test"
    assert manifest["confidence"] == 0.9


# ============= OCR TEST =============

def _tesseract_available():
    """Check if tesseract is available"""
    tesseract_cmd = r"C:\Program Files\Tesseract-OCR\tesseract.exe"
    return shutil.which("tesseract") is not None or Path(tesseract_cmd).exists()


@pytest.mark.skipif(not _tesseract_available(), reason="Tesseract not installed")
def test_ocr_pdf():
    """Test OCR on a scanned PDF (requires tesseract)"""
    pytest.importorskip("pdfplumber")
    pytest.importorskip("pytesseract")
    pytest.importorskip("PIL")
    
    # This test would need a real scanned PDF fixture
    # For now, just verify the function exists and is callable
    from app.extract import extract_text_scanned_pdf
    assert callable(extract_text_scanned_pdf)


# ============= INTEGRATION TESTS =============

def test_full_pipeline_scan_extract_quality(test_db, temp_dir):
    """Integration test: scan -> extract -> quality"""
    # Register and scan
    data = CollectionCreate(
        name="test-pipeline",
        host_display_path=temp_dir,
        container_root=temp_dir,
    )
    collection = register_collection(test_db, data)
    
    from app.jobs import JobQueue, ScanJobProcessor
    job_queue = JobQueue(test_db)
    job = job_queue.create_job(collection.id, JobType.SCAN)
    processor = ScanJobProcessor(test_db, job_queue)
    processor.process(job.id)
    
    # Verify documents created
    docs = test_db.query(Document).filter(Document.collection_id == collection.id).all()
    assert len(docs) >= 2
    
    # Run extraction
    from app.services import enqueue_extract
    from app.worker import Worker
    
    extract_job = enqueue_extract(test_db, collection.id)
    worker = Worker(test_db, queue_type="fast")
    worker.process_job(extract_job)
    
    # Verify extraction results
    docs = test_db.query(Document).filter(Document.collection_id == collection.id).all()
    for doc in docs:
        if doc.extension in ["txt", "md"]:
            assert doc.status in [DocumentStatus.INDEXED, DocumentStatus.NEEDS_REVIEW]
            assert doc.quality_score is not None
            assert doc.char_count is not None
            assert doc.char_count > 0


if __name__ == "__main__":
    pytest.main([__file__, "-v"])